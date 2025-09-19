import os
from zipfile import ZipFile
import shutil
import io
import PyPDF2
import pandas
from openpyxl import load_workbook

filenames = ['tmp/1mb.xlsx', 'tmp/Dummy-PDF-3Pages.pdf', 'tmp/example_2.5kb.csv']

# Создание и перемещение в папку ресурсы
with ZipFile('archive.zip', 'w') as my_zip:
    for filename in filenames:
        my_zip.write(filename)
os.makedirs('ресурсы', exist_ok=True)
destination_path = os.path.join('ресурсы', 'archive.zip')
shutil.move('archive.zip', destination_path)

# Проверка на то, что данный путь к файлу существует
if os.path.exists('ресурсы/archive.zip'):
    print('Данный путь к файлу существует')
else:
    print('Данный путь к файлу не существует')

# проверка pdf-файла
with ZipFile('ресурсы/archive.zip', 'r') as zip_file:
    print(zip_file.namelist())
    with zip_file.open('tmp/Dummy-PDF-3Pages.pdf') as pdf_file:
        reader = PyPDF2.PdfReader((io.BytesIO(pdf_file.read())))
        print(len(reader.pages))
        assert 'Sample T ext PDF' in reader.pages[0].extract_text()

    # проверка xlsx-файла
    with zip_file.open('tmp/1mb.xlsx') as xlsx_file:
        workbook = load_workbook((io.BytesIO(xlsx_file.read())))
        sheet = workbook.active
        print(sheet.cell(row=2, column=2).value)
        assert sheet.cell(row=2, column=2).value == 'rubye.bernhard@gmail.com'

    # проверка csv-файла
    with zip_file.open('tmp/example_2.5kb.csv') as csv_file:
        reader = pandas.read_csv(csv_file)['name'].tolist()
        print(reader)
        assert 'Prof. Gerardo Deckow V' in reader
